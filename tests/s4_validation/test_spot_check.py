"""E6 spot-check: draw, sheets, scoring."""

from __future__ import annotations

import numpy as np
import pandas as pd
import pytest

from skillsig.config import Config
from skillsig.s4_validation import spot_check
from skillsig.s4_validation.spot_check import (
    FORBIDDEN,
    SHEET_COLUMNS,
    STRATA,
    TAXONOMY,
    b_per_label,
    build_sheets,
    codebook_text,
    draw,
    quotas,
    score,
)


def _pool(per_stratum: int = 40) -> pd.DataFrame:
    rows = []
    i = 0
    for name in STRATA:
        for _ in range(per_stratum):
            rows.append(
                {
                    "file_sha": f"{i:040x}",
                    "stratum": name,
                    "repo_full_name": f"o{i % 300}/r{i}",
                    "cluster_id": float(i % 250),
                }
            )
            i += 1
    return pd.DataFrame(rows)


def test_draw_quotas_caps_and_determinism(cfg: Config) -> None:
    sample, design = draw(_pool(), cfg)
    assert len(sample) == 120
    assert sample.groupby("label_v21").size().to_dict() == {label: 30 for label in TAXONOMY}
    again, _ = draw(_pool(), cfg)
    assert sample.equals(again)
    assert sorted(sample["sample_id"])[0] == "S001"
    expected = quotas(cfg.sample.per_label)
    assert all(design[name]["drawn"] == expected[name] for name in STRATA)


def test_quotas_split_each_label_over_its_strata() -> None:
    assert quotas(15)["templated_registry"] == 8
    assert quotas(15)["templated_generic"] == 7
    assert set(quotas(30).values()) == {15}
    assert b_per_label(15) == 5
    assert b_per_label(30) == 10
    with pytest.raises(ValueError):
        quotas(0)


def test_draw_caps_one_per_repo_and_cluster(cfg: Config) -> None:
    pool = _pool()
    pool["repo_full_name"] = [f"o/r{i % 70}" for i in range(len(pool))]
    sample, _ = draw(pool, cfg)
    merged = sample.merge(pool, on=["file_sha", "stratum"])
    assert merged["repo_full_name"].is_unique
    assert merged["cluster_id"].is_unique


def test_shortfall_moves_to_sibling(cfg: Config) -> None:
    pool = _pool()
    pool = pool[~((pool["stratum"] == "human_revised") & (pool.index % 40 >= 3))]
    sample, design = draw(pool, cfg)
    assert design["human_revised"]["drawn"] == 3
    assert (sample["label_v21"] == "likely_human").sum() == 30


def _info(sample: pd.DataFrame) -> pd.DataFrame:
    n = len(sample)
    return pd.DataFrame(
        {
            "file_sha": sample["file_sha"],
            "repo_full_name": [f"o/r{i}" for i in range(n)],
            "skill_name": "demo",
            "description": "does things",
            "content": "---\nname: demo\n---\nBody text here.",
            "commit_count": 2,
            "first_commit_at": "2026-01-01",
            "last_commit_at": "2026-02-01",
            "first_commit_message": "add skill",
            "last_commit_message": "fix typo",
            "stars": 3,
            "language": "Python",
            "repo_created_at": "2025-12-01T00:00:00Z",
            "repo_files": 4,
            "label_v21": sample["label_v21"],
            "confidence_v21": "low",
            "reasons_v21": [["x"]] * n,
            "label_v2": sample["label_v21"],
            "reasons_v2": [["x"]] * n,
            "B0_majority": "insufficient_evidence",
        }
    )


def test_sheets_are_blind_and_b_is_subset(cfg: Config) -> None:
    sample, _ = draw(_pool(), cfg)
    sheets = build_sheets(sample, _info(sample), cfg)
    assert tuple(sheets["A"].columns) == SHEET_COLUMNS
    assert not FORBIDDEN & set(sheets["A"].columns)
    assert set(sheets["B"]["sample_id"]) <= set(sheets["A"]["sample_id"])
    assert len(sheets["B"]) == 40
    by_id = sheets["key"].set_index("sample_id")["file_sha"]
    assert (sheets["A"].set_index("sample_id")["file_sha"] == by_id).all()
    text = sheets["A"].astype(str).to_numpy().ravel()
    assert not any(label in cell for cell in text for label in TAXONOMY)
    assert not sheets["A"]["body_excerpt"].str.contains("name: demo").any()
    assert (sheets["A"]["body_chars"] == len("Body text here.")).all()


def test_excerpt_keeps_full_body_up_to_excel_limit() -> None:
    long = "---\nname: x\n---\n" + "w" * (spot_check.EXCERPT_CHARS + 500)
    assert spot_check._excerpt("---\nname: x\n---\nshort body") == "short body"
    assert len(spot_check._excerpt(long)) == spot_check.EXCERPT_CHARS
    assert len(spot_check._body(long)) == spot_check.EXCERPT_CHARS + 500


def test_score_precision_kappa_trailers(cfg: Config) -> None:
    key = pd.DataFrame(
        {
            "sample_id": ["S1", "S2", "S3", "S4"],
            "label_v21": ["agent_assisted", "agent_assisted", "likely_human", "highly_templated"],
            "label_v2": ["insufficient_evidence"] * 4,
            "reasons_v21": [["trailer_in_bulk"], ["coding_agent_bot"], ["revised"], ["copy_count"]],
            "reasons_v2": [["x"]] * 4,
            "B3_trailer": ["agent_assisted"] * 4,
        }
    )
    a = pd.DataFrame(
        {
            "sample_id": ["S1", "S2", "S3", "S4"],
            "your_label": [
                "agent_assisted",
                "insufficient_evidence",
                "likely_human",
                " Highly_Templated ",
            ],
        }
    )
    b = pd.DataFrame(
        {"sample_id": ["S1", "S2"], "your_label": ["agent_assisted", "insufficient_evidence"]}
    )
    trailer = pd.DataFrame({"is_ai_disclosure": ["yes", "yes", "no", "unsure", ""]})
    result = score(key, a, b, trailer)
    prec = result["precision"].set_index(["version", "label"])
    assert prec.loc[("label_v21", "agent_assisted"), "precision"] == pytest.approx(0.5)
    assert prec.loc[("label_v21", "highly_templated"), "precision"] == 1.0
    assert "indicative" in prec.loc[("label_v21", "likely_human"), "note"]
    reasons = result["reasons"].set_index(["version", "reason"])
    assert reasons.loc[("label_v21", "trailer_in_bulk"), "precision"] == 1.0
    assert result["agreement"]["kappa"] == 1.0
    assert result["trailers"]["precision"] == pytest.approx(2 / 3)


def test_run_without_done_file_draws_and_never_overwrites(tmp_cfg: Config, monkeypatch) -> None:
    calls = []
    monkeypatch.setattr(spot_check, "run_draw", lambda cfg, **kw: calls.append("draw"))
    spot_check.run(tmp_cfg)
    assert calls == ["draw"]
    sheet = spot_check._io.annotation(tmp_cfg, "spot_A.csv")
    sheet.parent.mkdir(parents=True)
    sheet.write_text("existing")
    monkeypatch.undo()
    assert spot_check.run_draw(tmp_cfg).skipped is True
    assert sheet.read_text() == "existing"


def test_score_requires_done_file(tmp_cfg: Config) -> None:
    with pytest.raises(FileNotFoundError, match="spot_A_done.xlsx"):
        spot_check.run_score(tmp_cfg)


def test_codebook_mentions_every_label() -> None:
    text = codebook_text(120, 40)
    assert "(120 rows)" in text and "(40 rows)" in text
    assert all(label in text for label in TAXONOMY)
    assert "not the true author" in text
    assert np.all([line == line.rstrip() for line in text.splitlines()])


def test_xlsx_sheet_has_dropdowns_and_reads_back(cfg: Config, tmp_path) -> None:
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    from skillsig.s4_validation import _sheets

    sample, _ = draw(_pool(), cfg)
    sheet = build_sheets(sample, _info(sample), cfg)["A"]
    sheet.loc[0, "first_commit_message"] = "=HYPERLINK(evil)\x07"
    path = tmp_path / "spot_A.xlsx"
    _sheets.write_annotation_xlsx(
        sheet,
        path,
        fill_columns=spot_check.FILL_COLUMNS,
        labels=TAXONOMY,
        instructions=codebook_text(120, 40),
        example=spot_check._EXAMPLE,
    )
    book = load_workbook(path)
    ws = book[_sheets.LABEL_SHEET]
    header = [c.value for c in ws[1]]
    assert header[0] == "sample_id"
    assert header[-3:] == ["your_label", "your_confidence", "notes"]
    label_col = get_column_letter(len(header) - 2)
    conf_col = get_column_letter(len(header) - 1)
    rules = {str(v.sqref): v.formula1 for v in ws.data_validations.dataValidation}
    assert rules == {
        f"{label_col}2:{label_col}{len(sheet) + 1}": '"' + ",".join(TAXONOMY) + '"',
        f"{conf_col}2:{conf_col}{len(sheet) + 1}": '"1,2,3,4,5"',
    }
    assert ws.auto_filter.ref is None
    assert all(c.fill.fgColor.rgb != "00FFFF00" for row in ws.iter_rows() for c in row)
    assert "Instructions" in book.sheetnames
    back = _sheets.read_annotation(path)
    assert len(back) == len(sheet)
    assert set(back.columns) == set(SHEET_COLUMNS)
    assert back.loc[0, "first_commit_message"] == "=HYPERLINK(evil)"  # text, not a formula


def test_score_from_returned_xlsx(tmp_cfg: Config) -> None:
    from openpyxl import load_workbook

    from skillsig.s4_validation import _sheets

    sample, _ = draw(_pool(), tmp_cfg)
    sheets = build_sheets(sample, _info(sample), tmp_cfg)
    done = spot_check._io.annotation(tmp_cfg, "spot_A_done.xlsx")
    _sheets.write_annotation_xlsx(
        sheets["A"],
        done,
        fill_columns=spot_check.FILL_COLUMNS,
        labels=TAXONOMY,
        instructions="",
        example={},
    )
    book = load_workbook(done)
    ws = book[_sheets.LABEL_SHEET]
    truth = dict(zip(sheets["key"]["sample_id"], sheets["key"]["label_v21"], strict=True))
    label_col = ws.max_column - 2
    for row in range(2, ws.max_row + 1):
        ws.cell(row=row, column=label_col, value=truth[ws.cell(row=row, column=1).value])
        ws.cell(row=row, column=label_col + 1, value=4)
    book.save(done)
    sheets["key"].to_parquet(spot_check._io.annotation(tmp_cfg, "key.parquet"), index=False)
    assert spot_check.done_path(tmp_cfg, "A") == done
    spot_check.run_score(tmp_cfg)
    prec = pd.read_csv(spot_check._io.results(tmp_cfg, "s4_spot_precision.csv"))
    v21 = prec[prec["version"] == "label_v21"].set_index("label")
    assert (v21["precision"] == 1.0).all()


def test_rebuild_keeps_sample_and_b_and_refreshes_labels(cfg: Config) -> None:
    sample, _ = draw(_pool(), cfg)
    info = _info(sample)
    old = build_sheets(sample, info, cfg)["key"]
    changed = info.copy()
    changed.loc[0, "label_v21"] = "likely_human"
    changed["content"] = "---\nname: demo\n---\nA much longer body now."
    rebuilt = spot_check.rebuild(old, changed, cfg)
    assert list(rebuilt["A"]["sample_id"]) == list(old["sample_id"])
    assert set(rebuilt["B"]["sample_id"]) == set(old.loc[old["in_b"], "sample_id"])
    key = rebuilt["key"].set_index("file_sha")
    first = changed.loc[0, "file_sha"]
    assert key.loc[first, "label_v21"] == "likely_human"
    assert key.loc[first, "label_v21_at_draw"] == old.set_index("file_sha").loc[first, "label_v21"]
    assert (rebuilt["A"]["body_excerpt"] == "A much longer body now.").all()
    assert set(rebuilt["key"]["stratum"]) == set(old["stratum"])


def test_spot_refresh_runner_resolves() -> None:
    from skillsig.cli import resolve_runner

    assert resolve_runner("sample", "spot_refresh") is spot_check.run_refresh
