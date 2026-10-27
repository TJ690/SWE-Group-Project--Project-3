"""
STAGE: s4_validation
MODULE: _sheets
PURPOSE: Write annotation sheets as Excel workbooks with dropdowns, and read them back.
INPUTS: none (callers pass the sheet frame and the allowed values)
OUTPUTS: none (writes the workbook path the caller passes)
RQ LINK: none (infrastructure for the E6 spot-check)

The workbook has a "Labels" sheet (one row per skill; the fill-in columns come last, and the
label and confidence columns only accept the allowed values) and an "Instructions" sheet (legend,
codebook, one example row that is not part of the sample). There is no header filter, so the only
dropdowns are the label and confidence lists.
"""

from __future__ import annotations

from collections.abc import Sequence
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

LABEL_SHEET = "Labels"
INSTRUCTIONS_SHEET = "Instructions"
CONFIDENCE_LEVELS = ("1", "2", "3", "4", "5")

_FONT = "Arial"
_HEADER_FILL = PatternFill("solid", start_color="D9D9D9")
_WIDE = {"body_excerpt": 90, "description": 45, "first_commit_message": 45}
_WIDE_DEFAULT = 30
_NARROW = {
    "sample_id": 10,
    "file_sha": 44,
    "your_label": 24,
    "your_confidence": 16,
    "commit_count": 12,
    "body_chars": 12,
}
_ROW_HEIGHT = 150


def write_annotation_xlsx(
    sheet: pd.DataFrame,
    path: Path,
    *,
    fill_columns: Sequence[str],
    labels: Sequence[str],
    instructions: str,
    example: dict[str, str],
) -> None:
    """Labels sheet (fill-in columns last, with dropdowns) plus an Instructions sheet."""
    rest = [c for c in sheet.columns if c not in fill_columns]
    columns = [*rest, *fill_columns]
    book = Workbook()
    labels_ws = book.active
    labels_ws.title = LABEL_SHEET
    _write_table(labels_ws, sheet[columns])
    _add_dropdowns(labels_ws, columns, len(sheet), labels)
    _write_instructions(book.create_sheet(INSTRUCTIONS_SHEET), instructions, example, columns)
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)


def read_annotation(path: Path) -> pd.DataFrame:
    """Labels sheet of a returned workbook, or a CSV with the same columns, as strings."""
    if path.suffix.lower() == ".xlsx":
        frame = pd.read_excel(path, sheet_name=LABEL_SHEET, dtype=str)
    else:
        frame = pd.read_csv(path, dtype=str)
    return frame.fillna("")


def _write_table(ws, frame: pd.DataFrame) -> None:
    header = Font(name=_FONT, bold=True)
    body = Font(name=_FONT)
    wrap = Alignment(wrap_text=True, vertical="top")
    for col, name in enumerate(frame.columns, start=1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.font = header
        cell.fill = _HEADER_FILL
        width = _NARROW.get(name, _WIDE.get(name, _WIDE_DEFAULT))
        ws.column_dimensions[get_column_letter(col)].width = width
    for row, values in enumerate(frame.itertuples(index=False), start=2):
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col)
            _set_text(cell, value)
            cell.font = body
            cell.alignment = wrap
        ws.row_dimensions[row].height = _ROW_HEIGHT
    ws.freeze_panes = "B2"  # header row and sample_id stay visible


def _add_dropdowns(ws, columns: Sequence[str], rows: int, labels: Sequence[str]) -> None:
    last = max(rows + 1, 2)
    rules = (
        ("your_label", labels, "Pick one of the four labels from the list."),
        ("your_confidence", CONFIDENCE_LEVELS, "Pick 1 (a guess) to 5 (you would bet on it)."),
    )
    for name, allowed, message in rules:
        if name not in columns:
            continue
        letter = get_column_letter(columns.index(name) + 1)
        rule = DataValidation(
            type="list",
            formula1='"' + ",".join(allowed) + '"',
            allow_blank=True,
            showErrorMessage=True,
            errorTitle="Not an allowed value",
            error=message,
            showInputMessage=True,
            promptTitle=name,
            prompt=message,
        )
        rule.add(f"{letter}2:{letter}{last}")
        ws.add_data_validation(rule)


def _write_instructions(ws, text: str, example: dict[str, str], columns: Sequence[str]) -> None:
    title = Font(name=_FONT, bold=True, size=13)
    body = Font(name=_FONT)
    bold = Font(name=_FONT, bold=True)
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 100
    ws.cell(row=1, column=1, value="How to fill in this workbook").font = title
    legend = [
        ("Read first", "docs/ANNOTATION_GUIDELINES.md: decision order, confidence, edge cases."),
        (
            "Edit",
            "Only the last three columns on the Labels sheet: your_label, your_confidence, notes.",
        ),
        ("your_label", "Dropdown: one of the four labels defined below."),
        (
            "your_confidence",
            "Dropdown: 1 = a guess, 2 = leaning, 3 = more likely than not, "
            "4 = fairly sure, 5 = you would bet on it.",
        ),
        ("notes", "Free text: what made you decide (optional but helpful)."),
        ("Do not", "Edit other columns, add or delete rows, look anything up, or discuss rows."),
    ]
    row = 3
    for key, value in legend:
        ws.cell(row=row, column=1, value=key).font = bold
        ws.cell(row=row, column=2, value=value).font = body
        row += 1
    row += 1
    ws.cell(row=row, column=1, value="Example (not part of the sample)").font = bold
    row += 1
    for name in ("sample_id", "your_label", "your_confidence", "notes"):
        if name in columns or name == "sample_id":
            ws.cell(row=row, column=1, value=name).font = bold
            cell = ws.cell(row=row, column=2)
            _set_text(cell, example.get(name, ""))
            cell.font = body
            row += 1
    row += 1
    ws.cell(row=row, column=1, value="Codebook").font = bold
    row += 1
    wrap = Alignment(wrap_text=True, vertical="top")
    for line in text.splitlines():
        cell = ws.cell(row=row, column=2)
        _set_text(cell, line.lstrip("# "))
        cell.font = bold if line.startswith("#") else body
        cell.alignment = wrap
        row += 1


def _set_text(cell, value: object) -> None:
    """Write as plain text: no formulas from '=...', no characters Excel rejects."""
    if value is None or (isinstance(value, float) and pd.isna(value)):
        cell.value = None
        return
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        cell.value = value
        return
    text = ILLEGAL_CHARACTERS_RE.sub("", str(value))[:32_000]
    cell.value = text
    cell.data_type = "s"
